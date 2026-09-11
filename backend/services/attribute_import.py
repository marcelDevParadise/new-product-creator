"""Preview and atomically apply product-specific XLSX attribute values."""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
import hashlib
import io
import json
import math
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook

from models.product import Product
from models.attribute import AttributeDefinition
from services import database
from services.shopify_metafields import normalize_value

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000


def text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def canonical(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)


def read_rows(content: bytes) -> list[dict]:
    if len(content) > MAX_FILE_BYTES:
        raise ValueError('Die XLSX-Datei darf höchstens 10 MB groß sein.')
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Die entpackte XLSX-Datei ist zu groß (maximal 50 MB).')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError('Die Datei ist keine lesbare XLSX-Datei.') from exc
    rows = []
    try:
        for sheet in workbook.worksheets:
            if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                raise ValueError(f'{sheet.title}: maximal {MAX_ROWS} Datenzeilen erlaubt.')
            if sheet.max_column and sheet.max_column > 100:
                raise ValueError(f'{sheet.title}: maximal 100 Spalten erlaubt.')
            iterator = sheet.iter_rows()
            header = next(iterator, None)
            if not header or not any(cell.value is not None for cell in header):
                continue
            headers = [text(cell.value).casefold() for cell in header]
            required = {'key', 'wert'}
            if not required.issubset(headers) or not {'produkt', 'artikelnummer', 'sku'}.intersection(headers):
                raise ValueError(f'{sheet.title}: benötigt werden Key, Wert und Produkt oder Artikelnummer/SKU.')
            nonempty = [h for h in headers if h]
            if len(nonempty) != len(set(nonempty)):
                raise ValueError(f'{sheet.title}: Spaltenüberschriften dürfen nicht doppelt vorkommen.')
            for index, cells in enumerate(iterator, 2):
                if not any(cell.value is not None for cell in cells):
                    continue
                row = {key: cell.value for key, cell in zip(headers, cells) if key}
                sku = text(row.get('artikelnummer') or row.get('sku'))
                source = sku or text(row.get('produkt'))
                rows.append({
                    'sheet': sheet.title, 'row': index, 'source': source,
                    'source_sku': sku, 'key': text(row.get('key')),
                    'value': row.get('wert'), 'declared_type': text(row.get('datentyp')),
                    'name': text(row.get('attributname')), 'category': text(row.get('kategorie')),
                    'evidence_kind': text(row.get('wertart')), 'source_url': text(row.get('quelle')),
                    'formula': any(cell.data_type in {'f', 'e'} for key, cell in zip(headers, cells)
                                   if key in {'produkt', 'artikelnummer', 'sku', 'key', 'wert', 'datentyp'}),
                })
                if len(rows) > MAX_ROWS:
                    raise ValueError(f'Maximal {MAX_ROWS} Datenzeilen pro Import erlaubt.')
    finally:
        workbook.close()
    if not rows:
        raise ValueError('Die Datei enthält keine Attributzeilen.')
    return rows


def convert_value(value, definition: AttributeDefinition):
    kind = definition.shopify_type
    if isinstance(value, (datetime, date)):
        value = value.date().isoformat() if kind == 'date' and isinstance(value, datetime) else value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        if kind.startswith('list.') or kind == 'json' or value.startswith('{'):
            try:
                value = json.loads(value)
            except json.JSONDecodeError:
                if kind.startswith('list.') or kind == 'json':
                    raise ValueError(f'{kind} benötigt gültiges JSON.')
        if kind == 'boolean' and isinstance(value, str):
            boolean = {'ja': True, 'nein': False, '1': True, '0': False, 'true': True, 'false': False}
            value = boolean.get(value.casefold(), value)
    if kind == 'boolean' and type(value) in {int, float} and value in (0, 1):
        value = bool(value)
    if kind == 'number_integer' and isinstance(value, float) and value.is_integer():
        value = int(value)
    if kind == 'date':
        date.fromisoformat(str(value))
    if kind == 'date_time':
        datetime.fromisoformat(str(value))
    if kind == 'single_line_text_field' and ('\n' in str(value) or '\r' in str(value)):
        raise ValueError('Einzeiliger Text darf keine Zeilenumbrüche enthalten.')
    result = normalize_value(value, kind, definition.unit)
    if isinstance(result, float) and not math.isfinite(result):
        raise ValueError('Der Wert muss eine endliche Zahl sein.')
    canonical(result)
    return result


def build_preview(content: bytes, products: dict[str, Product], definitions: dict[str, AttributeDefinition],
                  mapping: dict[str, str | None], mode: str = 'overwrite', target_sku: str | None = None) -> dict:
    if mode not in {'overwrite', 'fill_empty'}:
        raise ValueError('Unbekannter Importmodus.')
    if target_sku and target_sku not in products:
        raise ValueError('Das Zielprodukt wurde nicht gefunden.')
    rows = read_rows(content)
    grouped = defaultdict(list)
    for row in rows:
        grouped[row['source']].append(row)
    groups = []
    for source, items in grouped.items():
        candidates = []
        if items[0]['source_sku']:
            candidates = [p for p in products.values() if p.artikelnummer == source]
        else:
            eans = {text(r['value']) for r in items if r['key'].casefold() in {'ean', 'gtin'} and text(r['value'])}
            candidates = [p for p in products.values() if p.ean and p.ean in eans]
            if not candidates:
                candidates = [p for p in products.values() if p.artikelname.strip().casefold() == source.casefold()]
        suggested = candidates[0].artikelnummer if len(candidates) == 1 else None
        if target_sku:
            suggested = target_sku if suggested == target_sku or len(grouped) == 1 else None
        chosen = mapping.get(source, suggested)
        skipped = source in mapping and mapping[source] is None
        groups.append({'source': source, 'count': len(items), 'sku': chosen, 'skipped': skipped})
    targets = {g['source']: g for g in groups}
    result_rows = []
    updates = defaultdict(dict)
    seen = {}
    for row in rows:
        group = targets[row['source']]
        sku = group['sku']
        product = products.get(sku) if sku else None
        definition = definitions.get(row['key'])
        result = {**row, 'value': text(row['value']), 'sku': sku, 'old_value': None, 'status': 'error', 'message': ''}
        result.pop('formula')
        result.pop('source_sku')
        if group['skipped']:
            result.update(status='skipped', message='Produkt vom Import ausgeschlossen.')
        elif not row['source']:
            result['message'] = 'Produktzuordnung fehlt in der Datei.'
        elif not product:
            result['message'] = 'Bitte ein eindeutiges Zielprodukt auswählen.'
        elif target_sku and sku != target_sku:
            result['message'] = 'Hier können nur Attribute des geöffneten Produkts importiert werden.'
        elif row['formula']:
            result['message'] = 'Formeln und Excel-Fehler sind nicht erlaubt. Bitte feste Werte einfügen.'
        elif not definition:
            result['message'] = f'Unbekannter Attribut-Key: {row["key"] or "(leer)"}. Bitte zuerst im Attributkatalog anlegen.'
        elif row['declared_type'] and row['declared_type'] != definition.shopify_type:
            result['message'] = f'Datentyp passt nicht zum Katalog ({definition.shopify_type}).'
        elif row['value'] is None or text(row['value']) == '':
            result.update(status='skipped', message='Leerer Wert; vorhandener Wert bleibt erhalten.')
        else:
            try:
                value = convert_value(row['value'], definition)
                old = product.attributes.get(row['key'])
                result.update(value=value, old_value=old)
                pair = (sku, row['key'])
                if pair in seen:
                    if canonical(seen[pair]) != canonical(value):
                        raise ValueError('Widersprüchliche Werte für dasselbe Produkt und Attribut.')
                    result.update(status='skipped', message='Identische doppelte Zeile.')
                else:
                    seen[pair] = value
                    if canonical(old) == canonical(value):
                        result.update(status='unchanged', message='Bereits vorhanden.')
                    elif mode == 'fill_empty' and old not in (None, '', []):
                        result.update(status='skipped', message='Vorhandener Wert wird beibehalten.')
                    else:
                        result.update(status='update' if old is not None else 'add')
                        updates[sku][row['key']] = value
            except (ValueError, TypeError, OverflowError) as exc:
                result['message'] = str(exc)
        result_rows.append(result)
    preview = {
        'groups': groups, 'rows': result_rows, 'mode': mode,
        'errors': sum(r['status'] == 'error' for r in result_rows),
        'changes': sum(len(attrs) for attrs in updates.values()), 'products': len(updates),
        'skipped': sum(r['status'] == 'skipped' for r in result_rows),
        'unchanged': sum(r['status'] == 'unchanged' for r in result_rows),
    }
    preview['token'] = hashlib.sha256(canonical({
        'file': hashlib.sha256(content).hexdigest(), 'preview': preview,
        'definitions': {r['key']: definitions[r['key']].model_dump() for r in rows if r['key'] in definitions},
        'target': target_sku,
    }).encode()).hexdigest()
    return preview


def apply_preview(preview: dict, products: dict[str, Product], filename: str) -> None:
    """Commit all values/history together; reject concurrent changes before any write."""
    changes = [r for r in preview['rows'] if r['status'] in {'add', 'update'}]
    by_sku = defaultdict(list)
    for row in changes:
        by_sku[row['sku']].append(row)
    updated = {}
    with database.get_conn() as conn, conn.cursor() as cur:
        cur.execute('BEGIN')
        try:
            affected = set(by_sku)
            for sku in by_sku:
                product = products[sku]
                affected.add(product.parent_sku or sku)
                affected.update(p.artikelnummer for p in products.values() if p.parent_sku == sku)
            cur.execute("SELECT root_sku FROM articlewerk_jobs WHERE status IN ('queued', 'publishing')")
            if any(row[0] in affected for row in cur.fetchall()):
                raise ValueError('Für ein Zielprodukt läuft eine Übertragung. Bitte deren Ende abwarten.')
            for sku, rows in by_sku.items():
                cur.execute('SELECT attributes FROM products WHERE artikelnummer=%s', (sku,))
                record = cur.fetchone()
                if record is None:
                    raise ValueError('Ein Zielprodukt wurde entfernt. Bitte die Vorschau erneut laden.')
                raw = record[0]
                attributes = json.loads(raw)
                for row in rows:
                    if canonical(attributes.get(row['key'])) != canonical(row['old_value']):
                        raise ValueError('Attributwerte wurden inzwischen geändert. Bitte die Vorschau erneut laden.')
                    attributes[row['key']] = row['value']
                cur.execute('UPDATE products SET attributes=%s WHERE artikelnummer=%s AND attributes=%s',
                            (json.dumps(attributes, ensure_ascii=False), sku, raw))
                if cur.rowcount != 1:
                    raise ValueError('Produkt wurde inzwischen geändert. Bitte die Vorschau erneut laden.')
                updated[sku] = attributes
                for row in rows:
                    detail = {k: row[k] for k in ('sheet', 'row', 'source', 'source_url', 'evidence_kind')}
                    detail['file'] = filename
                    cur.execute(
                        f'INSERT INTO product_history (artikelnummer, event_type, field, old_value, new_value, detail, created_at) '
                        f"VALUES (%s, 'attribute_import', %s, %s, %s, %s, {database._NOW_SQL})",
                        (sku, row['key'], canonical(row['old_value']), canonical(row['value']), canonical(detail)),
                    )
            for sku in affected:
                cur.execute("UPDATE product_workflow SET status='in_progress', approved_hash=NULL, approved_at=NULL "
                            "WHERE artikelnummer=%s AND status='approved'", (sku,))
            cur.execute(f'INSERT INTO activity_log (event_type, detail, count, created_at) '
                        f"VALUES ('attribute_import', %s, %s, {database._NOW_SQL})",
                        (f'{len(changes)} Attributwerte aus {filename}', len(updated)))
            conn.commit()
        except Exception:
            conn.rollback()
            raise
    for sku, attributes in updated.items():
        products[sku].attributes = attributes
